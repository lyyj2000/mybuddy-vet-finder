#!/usr/bin/env python3
# ============================================================================
# build_map.py  —  MyBuddy Pet Finder (V30 - USER FEEDBACK REFINEMENTS)
# ============================================================================
import json, re, time, urllib.request, urllib.parse, random
from pathlib import Path
from openpyxl import load_workbook

WORKBOOK_VETS = "MyBuddy_Vet_Finder_MASTER_RECONCILED_2026-10-15_v9.xlsx"
WORKBOOK_PLACES = "PET_PLACES.xlsx"
WORKBOOK_REST = "Pets_places rest.xlsx"

OUTPUT   = "index.html"
CACHE_FILE = Path(__file__).parent / "geocache.json"
MANIFEST_OUTPUT = "manifest.json"
ICON_OUTPUT     = "icon.svg"

# SAFE FOR GITHUB UPLOAD (The GitHub Robot will inject the real key)
GOOGLE_API_KEY = "PASTE_YOUR_NEW_API_KEY_HERE"

for old in ["mybuddy-vet-map.html", "mybuddy-vet-map-v2.html", "mybuddy-vet-map-v3.html"]:
    old_file = Path(__file__).parent / old
    if old_file.exists():
        old_file.unlink()

DAY_NAMES = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"]
DAY_ALIASES = {"mon":"Monday","monday":"Monday","tue":"Tuesday","tues":"Tuesday","tuesday":"Tuesday","wed":"Wednesday","weds":"Wednesday","wednesday":"Wednesday","thu":"Thursday","thur":"Thursday","thurs":"Thursday","thursday":"Thursday","fri":"Friday","friday":"Friday","sat":"Saturday","saturday":"Saturday","sun":"Sunday","sunday":"Sunday"}

POSTAL_FALLBACKS = { "563338": "560338" }
EMERGENCY_OVERRIDES = { "MBV-0005", "MBV-0032" }

MANUAL_COORDS = { 
    "MBV-0005": (1.3244, 103.9418), "MBV-0032": (1.3286, 103.8521),
    "PLC-105": (1.3622, 103.8475), "PLC-106": (1.2901, 103.7622),
    "PLC-107": (1.3008, 103.9124), "PLC-108": (1.2860, 103.8236),
    "PLC-109": (1.4647, 103.8336), "PLC-110": (1.3115, 103.7955),
    "PLC-111": (1.3506, 103.8721), "PLC-112": (1.3065, 103.7885),
    "PLC-113": (1.2644, 103.8223), "PLC-114": (1.3054, 103.8294),
}

SPECIES_OVERRIDES = {}
SPECIES_ALIASES = {
    "dog":    ["dog", "canine"], "cat":    ["cat", "feline"],
    "bird":   ["bird", "avian", "parrot"],
    "exotic": ["exotic", "rabbit", "small", "reptile", "pocket", "hamster", "guinea"],
}

def to_bool(v):
    if isinstance(v, bool): return v
    if isinstance(v, str):  return v.strip().lower() in ("true","yes","1","y")
    return False

def parse_time(s):
    s = str(s).strip().lower().replace(".", "").replace(" ", "")
    if not s: return None
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?\s*(am|pm)?$", s)
    if not m: return None
    h = int(m.group(1)); mi = int(m.group(2) or 0); ap = m.group(3)
    if mi > 59: return None
    if ap == "am" and h == 12: h = 0
    elif ap == "pm" and h != 12: h += 12
    elif not ap and h > 23: return None
    return h * 60 + mi

def parse_intervals(text):
    if text is None: return None
    t = str(text).strip().lower()
    if not t: return None
    if re.search(r"\bby\s+appointment\b", t): return "appointment"
    if re.match(r"^closed\b", t): return []
    if re.search(r"\b(open\s+)?24\s*hours?\b", t) or "24/7" in t: return [(0, 1440)]
    ranges = re.findall(r"(\d{1,2}(?::\d{2})?\s*(?:am|pm)?)\s*(?:[\u2013\u2014\-]|to)\s*(\d{1,2}(?::\d{2})?\s*(?:am|pm)?)", t, re.I)
    if not ranges: return [] if re.search(r"\bclosed\b", t) else None
    intervals = []
    for a, b in ranges:
        sa = parse_time(a); sb = parse_time(b)
        if sa is None or sb is None: continue
        if sb <= sa: sb += 1440
        intervals.append((sa, sb))
    return intervals if intervals else None

def day_in_range(start, end, target):
    s = DAY_ALIASES.get(start.lower()); e = DAY_ALIASES.get(end.lower())
    if not s or not e: return False
    si = DAY_NAMES.index(s); ei = DAY_NAMES.index(e); cur = si
    while True:
        if DAY_NAMES[cur] == target: return True
        if cur == ei: break
        cur = (cur + 1) % 7
    return False

def extract_day_schedule(full_text, target_day):
    if not full_text: return ""
    chunks = re.split(r"[;\n]+", str(full_text))
    result = []; current_days = None
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk: continue
        dr = re.match(r"^(mon|monday|tue|tues|tuesday|wed|weds|wednesday|thu|thur|thurs|thursday|fri|friday|sat|saturday|sun|sunday)\s*[\u2013\u2014\-]\s*(mon|monday|tue|tues|tuesday|wed|weds|wednesday|thu|thur|thurs|thursday|fri|friday|sat|saturday|sun|sunday)\b", chunk, re.I)
        if dr:
            current_days = [d for d in DAY_NAMES if day_in_range(dr.group(1), dr.group(2), d)]
            if target_day in current_days:
                rest = chunk[dr.end():].strip()
                if rest: result.append(rest)
            continue
        ds = re.match(r"^(mon|monday|tue|tues|tuesday|wed|weds|wednesday|thu|thur|thurs|thursday|fri|friday|sat|saturday|sun|sunday)\b", chunk, re.I)
        if ds:
            d = DAY_ALIASES.get(ds.group(1).lower())
            current_days = [d]
            if d == target_day:
                rest = chunk[ds.end():].strip()
                if rest: result.append(rest)
            continue
        if current_days is None or target_day in current_days: result.append(chunk)
    return "; ".join(result)

def read_sheet(wb, name_or_index):
    try:
        if isinstance(name_or_index, int): ws = wb.worksheets[name_or_index]
        else: ws = wb[name_or_index]
    except KeyError: return []
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    return [{headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))} for row in rows[1:] if not all(c is None for c in row)]

def _species_from_text(s):
    s = str(s or "").strip().lower()
    if not s: return set()
    return {tag for tag, kws in SPECIES_ALIASES.items() if any(kw in s for kw in kws)}

def extract_species_tags(services_data):
    tags = {}
    if not services_data: return tags
    headers = list(services_data[0].keys())
    has_species_col = any(str(h).strip().lower() == "species" for h in headers)
    wide_cols = [(h, _species_from_text(h)) for h in headers if _species_from_text(h)]
    for r in services_data:
        mbv = str(r.get("MyBuddy Vet ID") or "").strip()
        if not mbv: continue
        if has_species_col:
            sp_set = _species_from_text(r.get("Species"))
            if sp_set: tags.setdefault(mbv, set()).update(sp_set)
        else:
            for col, sp_set in wide_cols:
                if to_bool(r.get(col)): tags.setdefault(mbv, set()).update(sp_set)
    return {k: sorted(v) for k, v in tags.items()}

def _query_onemap(q):
    url = "https://www.onemap.gov.sg/api/common/elastic/search?" + urllib.parse.urlencode({
        "searchVal": q, "returnGeom": "Y", "getAddrDetails": "Y", "pageNum": 1})
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=8) as r:
            data = json.loads(r.read().decode("utf-8"))
        for p in (data.get("results") or []):
            try:
                lat = float(p["LATITUDE"]); lng = float(p["LONGITUDE"])
                if 1.1 <= lat <= 1.5 and 103.5 <= lng <= 104.2: return (lat, lng)
            except (KeyError, ValueError): continue
    except Exception: pass
    return None

def geocode_location(postal, address, name, cache):
    key = f"{postal}|{address}|{name}"
    if key in cache: return cache[key]
    queries = []
    if postal and len(str(postal).strip()) == 6: queries.append(str(postal).strip())
    if name: queries.append(f"{name} Singapore")
    if address:
        clean_addr = str(address).replace('\n', ' ').strip()
        queries.append(f"{clean_addr} Singapore")
    for q in queries:
        coords = _query_onemap(q)
        if coords:
            cache[key] = coords
            return coords
    return None

def fetch_google_data(name, address, postal):
    if GOOGLE_API_KEY == "PASTE_YOUR_NEW_API_KEY_HERE": return None, None, None, None, None
    queries = [f"{name} {postal} Singapore", f"{name} Singapore", f"{postal} Singapore", f"{address} Singapore"]
    for q in queries:
        if not q.strip() or q.strip() == "Singapore": continue
        url = f"https://maps.googleapis.com/maps/api/place/textsearch/json?query={urllib.parse.quote(q)}&key={GOOGLE_API_KEY}"
        try:
            req = urllib.request.Request(url)
            with urllib.request.urlopen(req, timeout=5) as r:
                data = json.loads(r.read().decode("utf-8"))
                if data.get("results"):
                    first = data["results"][0]
                    loc = first.get("geometry", {}).get("location", {})
                    if loc.get("lat") and loc.get("lng"):
                        return (first.get("rating", 0), first.get("user_ratings_total", 0), 
                                loc.get("lat"), loc.get("lng"), first.get("place_id", ""))
        except Exception: pass
    return None, None, None, None, None

def write_pwa_files():
    manifest = {
        "name": "MyBuddy Pet Finder", "short_name": "Pet Finder",
        "start_url": "./index.html", "scope": "./",
        "display": "standalone", "orientation": "portrait",
        "background_color": "#ffffff", "theme_color": "#4CAF50",
        "icons": [{"src": "icon.svg", "sizes": "any", "type": "image/svg+xml", "purpose": "any"}]
    }
    Path(MANIFEST_OUTPUT).write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    icon_svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 512 512"><rect width="512" height="512" rx="96" fill="#4CAF50"/><circle cx="176" cy="180" r="44" fill="#ffffff"/><circle cx="256" cy="140" r="44" fill="#ffffff"/><circle cx="336" cy="180" r="44" fill="#ffffff"/><circle cx="128" cy="272" r="40" fill="#ffffff"/><circle cx="384" cy="272" r="40" fill="#ffffff"/><path d="M256 240c-58 0-104 46-104 104 0 46 38 68 104 68s104-22 104-68c0-58-46-104-104-104z" fill="#ffffff"/></svg>'
    Path(ICON_OUTPUT).write_text(icon_svg, encoding="utf-8")

def build():
    wb_vets = load_workbook(WORKBOOK_VETS, data_only=True)
    master = read_sheet(wb_vets, "MASTER_LIST")
    hours = read_sheet(wb_vets, "VET_CONTACT_HOURS")
    branches = read_sheet(wb_vets, "BRANCH_RELATIONSHIPS")
    prices_data = read_sheet(wb_vets, "VET_PRICES")
    services_data = read_sheet(wb_vets, "VET_SERVICES")
    
    places_data = []
    try:
        wb_places = load_workbook(WORKBOOK_PLACES, data_only=True)
        places_data.extend(read_sheet(wb_places, 0))
    except Exception as e: print(f"Warning: Could not load {WORKBOOK_PLACES} - {e}")
    try:
        wb_rest = load_workbook(WORKBOOK_REST, data_only=True)
        places_data.extend(read_sheet(wb_rest, 0))
    except Exception as e: print(f"Warning: Could not load {WORKBOOK_REST} - {e}")

    hours_by_id = {r["MyBuddy Vet ID"]: r for r in hours if r.get("MyBuddy Vet ID")}
    prices_by_id = {}
    for p in prices_data:
        mbv = str(p.get("MyBuddy Vet ID") or "").strip()
        if not mbv: continue
        service = p.get("Service"); price = p.get("Price")
        if service and price not in (None, ""): prices_by_id.setdefault(mbv, []).append(f"{service}: ${price}")

    species_by_id = extract_species_tags(services_data)
    for _id, _sp in SPECIES_OVERRIDES.items():
        species_by_id[_id] = sorted(set(species_by_id.get(_id, [])) | set(_sp))

    groups = {}
    for r in branches:
        gid = r.get("Group ID")
        if not gid: continue
        groups.setdefault(gid, []).append({"id": r.get("MBV ID"), "name": r.get("Branch Name") or ""})
    siblings = {}
    for gid, members in groups.items():
        for m in members: siblings[m["id"]] = [x for x in members if x["id"] != m["id"]]

    clinics = []
    
    # 1. PROCESS VETS
    for r in master:
        if (r.get("Active / Closed / Relocated Status") or "").strip() != "Current": continue
        mbv = str(r.get("MyBuddy Vet ID") or "").strip()
        h = hours_by_id.get(mbv, {})
        postal = str(int(r.get("Postal Code"))) if isinstance(r.get("Postal Code"), (int, float)) else str(r.get("Postal Code") or "").strip()
        if postal.endswith(".0"): postal = postal[:-2]
        postal = postal.zfill(6)
        addr = ", ".join(str(p).strip() for p in [r.get("Street Address"), r.get("Building Name")] if p)
        sched = {day: parse_intervals(extract_day_schedule(h.get(day), day)) if h.get(day) else None for day in DAY_NAMES}
        ecat = (h.get("Emergency Category") or "NOT_24_7_UNCONFIRMED").strip()
        if mbv in EMERGENCY_OVERRIDES: ecat = "E1_CONFIRMED_24_7"
        appt_req = str(h.get("Appointment Required") or "").strip()
        is_appt = "appointment" in appt_req.lower() and "no" not in appt_req.lower()
        clinics.append({
            "id": mbv, "category": "Vet", "name": str(r.get("Clinic / Provider Name") or ""),
            "branch": r.get("Branch Name") or "", "address": addr, "postal": postal,
            "phone": (r.get("Phone / Main Contact") or h.get("Main Phone") or "").strip(),
            "whatsapp": (r.get("WhatsApp") or h.get("WhatsApp") or "").strip(),
            "booking": h.get("Booking URL") or "", "ecat": ecat, "schedule": sched,
            "house_call": to_bool(r.get("Offers_House_Calls")), "hospice": to_bool(r.get("Hospice_Service")),
            "lat": None, "lng": None, "rating": r.get("Google Rating") or 0,
            "review_count": r.get("Review Count") or 0, "is_appt": is_appt,
            "prices": prices_by_id.get(mbv, []), "species": species_by_id.get(mbv, []),
            "hours_text": "", "place_id": ""
        })

    # 2. PROCESS PLACES
    for p in places_data:
        pid = str(p.get("Place ID") or p.get("ID") or "").strip()
        if not pid: continue
        raw_cat = str(p.get("Category") or "").strip().lower()
        if "park" in raw_cat: cat = "Park"
        elif "mall" in raw_cat: cat = "Mall"
        elif "groom" in raw_cat: cat = "Groomer"
        elif "restaurant" in raw_cat or "dining" in raw_cat or "food" in raw_cat or "cafe" in raw_cat: cat = "Restaurant"
        else: cat = raw_cat.capitalize() or "Other"
        rating_val = p.get("Google Rating") or p.get("Rating")
        try: rating_val = float(rating_val) if rating_val else 0.0
        except: rating_val = 0.0
        postal = str(int(p.get("Postal Code"))) if isinstance(p.get("Postal Code"), (int, float)) else str(p.get("Postal Code") or "").strip()
        if postal.endswith(".0"): postal = postal[:-2]
        postal = postal.zfill(6)
        lat_val = p.get("Latitude")
        lng_val = p.get("Longitude")
        try:
            if lat_val and lng_val: lat_val = float(lat_val); lng_val = float(lng_val)
            else: lat_val = None; lng_val = None
        except: lat_val = None; lng_val = None
        clinics.append({
            "id": pid, "category": cat, "name": str(p.get("Name") or "").strip(),
            "branch": "", "address": p.get("Address") or "", "postal": postal,
            "phone": str(p.get("Phone") or "").strip(), "whatsapp": "", "booking": "",
            "ecat": "", "schedule": {}, "house_call": False, "hospice": False,
            "lat": lat_val, "lng": lng_val, "rating": rating_val, "review_count": 0, "is_appt": False,
            "prices": [], "species": [], "hours_text": str(p.get("Operating Hours") or "").strip(),
            "place_id": ""
        })

    cache = {}
    if CACHE_FILE.exists():
        try: cache = {k: tuple(v) for k, v in json.loads(CACHE_FILE.read_text()).items()}
        except Exception: pass

    for c in clinics:
        if c["id"] in MANUAL_COORDS:
            c["lat"], c["lng"] = MANUAL_COORDS[c["id"]]
            if c["category"] in ["Vet", "Restaurant"]:
                live_rating, live_reviews, _, _, place_id = fetch_google_data(c['name'], c['address'], c['postal'])
                if live_rating is not None:
                    c["rating"] = live_rating; c["review_count"] = live_reviews or 0
                    c["place_id"] = place_id or ""
        else:
            live_rating, live_reviews, g_lat, g_lng, place_id = fetch_google_data(c['name'], c['address'], c['postal'])
            if live_rating is not None:
                c["rating"] = live_rating; c["review_count"] = live_reviews or 0
                c["place_id"] = place_id or ""
            if g_lat and g_lng:
                c["lat"], c["lng"] = g_lat, g_lng
            if c["lat"] is None or c["lng"] is None:
                coords = geocode_location(c["postal"], c["address"], c["name"], cache)
                if coords: c["lat"], c["lng"] = coords
                else: c["lat"] = None; c["lng"] = None
        time.sleep(0.05)

    try: CACHE_FILE.write_text(json.dumps({k: v for k, v in cache.items()}, indent=2))
    except Exception: pass

    valid_clinics = [c for c in clinics if c["lat"] is not None and c["lng"] is not None]
    seen_coords = set()
    for c in valid_clinics:
        coord_key = (round(c["lat"], 5), round(c["lng"], 5))
        if coord_key in seen_coords:
            c["lat"] += random.uniform(-0.0001, 0.0001); c["lng"] += random.uniform(-0.0001, 0.0001)
        else: seen_coords.add(coord_key)

    directory = {}
    for c in valid_clinics:
        letter = c["name"][:1].upper()
        if not letter.isalpha(): letter = "#"
        directory.setdefault(letter, []).append({"id": c["id"], "name": c["name"], "category": c["category"]})
    for L in directory: directory[L].sort(key=lambda x: x["name"].lower())

    data_json = json.dumps({"clinics": valid_clinics, "siblings": siblings, "directory": directory}, ensure_ascii=False, separators=(",", ":"))
    Path(OUTPUT).write_text(HTML_TEMPLATE.replace("/*DATA*/", data_json), encoding="utf-8")
    write_pwa_files()
    print(f"SUCCESS: Wrote {OUTPUT} with {len(valid_clinics)} mapped locations.")

HTML_TEMPLATE = r'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover,maximum-scale=1">
<title>MyBuddy Pet Finder</title>
<link rel="manifest" href="manifest.json">
<meta name="theme-color" content="#4CAF50">
<link rel="icon" type="image/svg+xml" href="icon.svg">
<link href="https://fonts.googleapis.com/css2?family=Baloo+2:wght@500;600;700;800&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
<style>
html,body,#map{height:100%;margin:0;font-family:'Baloo 2', cursive}
#map{background:#e8edf0}
#panel{position:absolute;z-index:1000;top:12px;left:12px;width:min(360px,calc(100vw - 40px));max-height:calc(100vh - 40px);overflow:visible;background:#fff;padding:16px;border-radius:20px;box-shadow:0 8px 25px rgba(0,0,0,.1);font-size:14px;display:flex;flex-direction:column;}
.panel-content { overflow-y: auto; flex: 1; padding-right: 4px; }
h1{font-size:22px;margin:0 0 4px;color:#388E3C;font-weight:800}
p.sub{font-size:12px;color:#90A4AE;margin:0 0 10px;font-weight:600}
.row{display:flex;align-items:center;gap:8px;margin:10px 0}
.row label{font-weight:700;color:#2C3E50;white-space:nowrap;font-size:13px}
#radius{flex:1;min-width:80px;accent-color:#4CAF50}
#radiusVal{min-width:52px;text-align:right;color:#4CAF50;font-weight:800}
.btn{background:#4CAF50;color:#fff;border:none;padding:10px 14px;border-radius:14px;font-size:13px;font-family:inherit;font-weight:700;cursor:pointer;flex:1;transition:transform 0.2s, background 0.2s;}
.btn:hover{background:#388E3C;transform:scale(1.02)}
.btn.ghost{background:#F1F8E9;color:#388E3C;border:1px solid #C8E6C9}
.btn.ghost:hover{background:#E8F5E9}
.input{width:100%;box-sizing:border-box;padding:10px 14px;border:2px solid #E0E0E0;border-radius:14px;font-size:13px;font-family:inherit;font-weight:600}
.input:focus{outline:none;border-color:#4CAF50}
.search-wrap{position:relative;margin:10px 0}
.search-wrap .input{padding-right:30px}
.search-wrap::after{content:"\1F50D";position:absolute;right:12px;top:50%;transform:translateY(-50%);font-size:14px;color:#90A4AE;pointer-events:none}
.directory{display:none;position:absolute;top:100%;left:0;right:0;background:#fff;border:2px solid #E0E0E0;border-radius:14px;margin-top:6px;max-height:300px;overflow-y:auto;box-shadow:0 12px 30px rgba(0,0,0,.15);z-index:1001}
.directory.open{display:block}
.dir-letter{border-bottom:1px solid #F0F0F0}
.dir-header{padding:10px 16px;font-weight:800;color:#388E3C;cursor:pointer;background:#F9FBF9;display:flex;justify-content:space-between}
.dir-header::after{content:"+";font-weight:600;color:#90A4AE;font-size:16px}
.dir-letter.open .dir-header::after{content:"−"}
.dir-items{display:none;padding:4px 0}
.dir-letter.open .dir-items{display:block}
.dir-item{padding:10px 16px 10px 24px;cursor:pointer;font-size:13px;font-weight:600;color:#2C3E50}
.dir-item:hover{background:#E8F5E9;color:#388E3C}
.filters{display:flex;flex-wrap:wrap;gap:6px;margin:10px 0}
.chip{font-size:12px;padding:8px 14px;border-radius:99px;border:2px solid #E0E0E0;background:#fff;cursor:pointer;font-weight:700;color:#90A4AE;transition:all 0.2s}
.chip.on{background:#4CAF50;color:#fff;border-color:#4CAF50}
.legend{font-size:12px;line-height:1.8;color:#90A4AE;border-top:2px dashed #E0E0E0;padding-top:12px;margin-top:12px;font-weight:600}
.dot{display:inline-block;width:12px;height:12px;border-radius:50%;margin-right:6px;vertical-align:middle;border:2px solid #fff;box-shadow:0 1px 3px rgba(0,0,0,.2)}
#status{font-size:12px;color:#388E3C;margin-top:10px;font-weight:700;background:#E8F5E9;padding:8px 12px;border-radius:10px;display:inline-block}
.disclaimer { font-size: 11px; color: #90A4AE; margin-top: 16px; padding-top: 12px; border-top: 2px dashed #E0E0E0; font-weight: 600; line-height: 1.4; }
.disclaimer b { color: #388E3C; }

.pin{width:22px;height:22px;border-radius:50%;border:3px solid #fff;box-shadow:0 2px 8px rgba(0,0,0,.3);display:flex;align-items:center;justify-content:center;font-size:12px;font-weight:bold;}
.pin.open{background:#4CAF50; color:#fff;}
.pin.closed{background:#90A4AE; color:#fff;}
.pin.emergency{background:#F44336; color:#fff; font-size:16px; animation:pulse 1.5s infinite;}
.pin.groomer{background:#FFEB3B; color:#000; border-color:#fff;}
.pin.park{background:#8BC34A; color:#fff; animation:pulse-brown 1.5s infinite;}
.pin.mall{background:#FFB300; color:#fff;}
.pin.dining{background:#fff; border:3px solid #000; color:#000; box-shadow: 0 0 10px rgba(0,0,0,0.2);}

@keyframes pulse{0%{box-shadow:0 0 0 0 rgba(244,67,54,0.6)}70%{box-shadow:0 0 0 12px rgba(244,67,54,0)}100%{box-shadow:0 0 0 0 rgba(244,67,54,0)}}
@keyframes pulse-brown{0%{box-shadow:0 0 0 0 rgba(139,69,19,0.6)}70%{box-shadow:0 0 0 12px rgba(139,69,19,0)}100%{box-shadow:0 0 0 0 rgba(139,69,19,0)}}

.pin-hit{width:36px;height:36px;display:flex;align-items:center;justify-content:center;cursor:pointer}
.centre-handle-icon{background:transparent!important;border:none!important}
.centre-handle{width:46px;height:46px;border:3px solid rgba(33,150,243,.5);border-radius:50%;background:rgba(33,150,243,.1);display:flex;align-items:center;justify-content:center;color:#2196F3;font-size:30px;font-weight:bold;cursor:grab}

.popup{font-size:13px;line-height:1.5;min-width:260px;max-width:300px;font-family:'Baloo 2', cursive}
.popup h3{font-size:16px;margin:0 0 4px;color:#2C3E50;font-weight:800}
.popup p{margin:4px 0;font-weight:600}
.badge{display:inline-block;font-size:11px;padding:3px 8px;border-radius:99px;margin:0 4px 4px 0;color:#fff;font-weight:700}
.leaflet-popup-content .btn-link { display: inline-block; margin: 6px 6px 0 0; padding: 8px 12px; border-radius: 10px; text-decoration: none; color: #FFEB3B !important; text-shadow: 0 1px 2px rgba(0,0,0,0.4); background: #2196F3; font-size: 13px; font-weight: 800; text-align: center; }
.leaflet-popup-content .btn-link.wa { background: #4CAF50; }
.price-list { max-height: 120px; overflow-y: auto; background: #F9FBFD; padding: 10px; border-radius: 10px; font-size: 12px; border: 1px solid #E1E8ED; margin-top: 8px; }
.price-item { padding: 4px 0; border-bottom: 1px dashed #E0E0E0; font-weight: 600; color: #555; }
.price-item:last-child { border-bottom: none; }

.slider-container { display: flex; align-items: center; gap: 10px; margin-top: 6px; background: #fff; border: 2px solid #E0E0E0; padding: 8px 12px; border-radius: 14px; }
.slider-container input[type=range] { flex: 1; accent-color: #FF9800; cursor: pointer; }
.slider-val { font-size: 12px; font-weight: 800; color: #FF9800; min-width: 70px; text-align: right; }

#mobile-top-bar, #mobile-species-bar, #mobile-radius-slider, #mobile-rating-bar { display: none; }
@media (max-width: 768px) {
  #panel { display: none !important; }
  #mobile-top-bar { display: flex; position: fixed; top: calc(env(safe-area-inset-top, 0px) + 12px); left: 12px; right: 12px; z-index: 2500; gap: 8px; }
  .m-search-wrap { flex: 1; position: relative; }
  #mobile-search-input { width: 100%; height: 46px; padding: 0 16px; border: none; border-radius: 16px; box-shadow: 0 4px 15px rgba(0,0,0,.15); font-family: inherit; font-size: 15px; font-weight: 600; background: #fff; color: #2C3E50; box-sizing: border-box; }
  
  #mobile-species-bar { display: flex; position: fixed; top: calc(env(safe-area-inset-top, 0px) + 68px); left: 12px; right: 12px; z-index: 2500; gap: 6px; white-space: nowrap; overflow-x: auto; padding-bottom: 4px; -webkit-overflow-scrolling: touch; }
  #mobile-species-bar::-webkit-scrollbar { display: none; }
  .m-sp-btn { flex: 0 0 auto; height: 40px; padding: 0 16px; border: none; border-radius: 12px; background: #fff; box-shadow: 0 4px 15px rgba(0,0,0,.12); font-size: 14px; font-weight: 700; cursor: pointer; transition: background 0.2s; }
  .m-sp-btn.on { background: #4CAF50; color: #fff; }
  
  #mobile-rating-bar { display: none; position: fixed; top: calc(env(safe-area-inset-top, 0px) + 116px); left: 12px; right: 12px; z-index: 2500; background: #fff; padding: 8px 12px; border-radius: 12px; box-shadow: 0 4px 15px rgba(0,0,0,.15); align-items: center; gap: 10px; }
  #mobile-rating-bar.active { display: flex; }
  #mobile-rating-bar input { flex: 1; accent-color: #FF9800; }
  #mobile-rating-val { font-size: 13px; font-weight: 800; color: #FF9800; min-width: 60px; text-align: right; }
  
  #mobile-radius-slider { display: block; position: fixed; bottom: calc(env(safe-area-inset-bottom, 0px) + 20px); left: 12px; right: 12px; z-index: 2500; background: rgba(255,255,255,0.95); border-radius: 16px; box-shadow: 0 4px 15px rgba(0,0,0,.2); padding: 10px 16px; }
  #mobile-radius-slider input { width: 100%; accent-color: #4CAF50; }

  #mobile-dir-dropdown { display: none; position: absolute; top: 52px; left: 0; right: 0; background: #fff; border-radius: 16px; box-shadow: 0 8px 25px rgba(0,0,0,.2); max-height: 65vh; overflow-y: auto; z-index: 2600; }
  #mobile-dir-dropdown.open { display: block; }
  .md-letter { padding: 12px 16px; font-weight: 800; font-size: 15px; color: #388E3C; background: #F9FBF9; border-bottom: 1px solid #EEE; cursor: pointer; display: flex; justify-content: space-between; align-items: center; }
  .md-letter::after { content: "+"; font-size: 18px; color: #90A4AE; }
  .md-letter.open::after { content: "−"; }
  .md-items { display: none; }
  .md-letter.open + .md-items { display: block; }
  .md-item { padding: 14px 16px 14px 24px; border-bottom: 1px solid #F1F5F9; font-weight: 600; font-size: 14px; color: #2C3E50; cursor: pointer; }
  .md-item:active { background: #E8F5E9; }

  .leaflet-control-zoom { display: none !important; }
  .leaflet-popup-content-wrapper { border-radius: 20px !important; }
  .leaflet-popup-content { margin: 14px 16px !important; }
  .popup { max-width: calc(100vw - 60px) !important; }
}
.leaflet-tile-pane { filter: saturate(1.5) contrast(0.95) hue-rotate(-5deg) brightness(1.05); }
</style>
</head>
<body>
<div id="map"></div>

<div id="mobile-top-bar">
  <div class="m-search-wrap">
    <input id="mobile-search-input" placeholder="Search name or postal code..." autocomplete="off">
    <div id="mobile-dir-dropdown"></div>
  </div>
</div>

<div id="mobile-species-bar">
  <button class="m-sp-btn cat-btn on" data-cat="Vet">🏥 Vets</button>
  <button class="m-sp-btn cat-btn" data-cat="Groomer">🌼 Groomers</button>
  <button class="m-sp-btn cat-btn" data-cat="Restaurant">🍴 Dine with Pets</button>
  <button class="m-sp-btn cat-btn" data-cat="Mall">🛍️ Pet Mall</button>
</div>

<div id="mobile-rating-bar">
  <input type="range" id="mobile-rating-slider" min="1" max="5" step="0.1" value="1">
  <span id="mobile-rating-val">All Stars</span>
</div>

<div id="mobile-radius-slider">
  <input id="mobile-radius-input" type="range" min="1" max="30" value="2" step="1">
</div>

<section id="panel">
  <div class="panel-content">
      <h1>🐾 MyBuddy Pet Finder</h1>
      
      <div class="filters" id="cat-filters" style="margin-top:15px;">
        <button class="chip cat-btn on" data-cat="Vet" style="width: 100%; font-size: 15px; padding: 12px; border-radius: 12px; background:#4CAF50; color:#fff; border-color:#4CAF50;">🏥 Vet Finder</button>
      </div>

      <div class="filters" id="species-filters" style="margin-top: 4px;">
        <button class="chip sp-btn" data-sp="dog" style="flex: 1;">🐶 Dogs</button>
        <button class="chip sp-btn" data-sp="cat" style="flex: 1;">🐱 Cats</button>
        <button class="chip sp-btn" data-sp="exotic" style="flex: 1;">🐰 Exotics</button>
      </div>

      <div class="filters" style="margin-top: 12px;">
        <button class="chip cat-btn" data-cat="Groomer" style="width: 100%; font-size: 14px; padding: 10px;">🌼 Groomers</button>
      </div>

      <div class="filters" style="margin-top: 6px;">
        <button class="chip cat-btn" data-cat="Mall" style="width: 100%; font-size: 14px; padding: 10px;">🛍️ Pet Mall</button>
      </div>

      <div class="search-wrap" style="margin-top: 16px;">
        <input id="search" class="input" placeholder="Search name or postal code...">
        <div id="directory" class="directory"></div>
      </div>

      <div class="filters" style="margin-top: 4px; flex-direction: column; gap: 8px;">
        <button class="chip cat-btn" data-cat="Restaurant" style="width: 100%; font-size: 14px; padding: 10px;">🍴 Dine with Pets</button>
        <div id="desktop-rating-container" class="slider-container" style="display:none;">
            <span style="font-size:14px; color:#FF9800;">⭐</span>
            <input type="range" id="rating-slider" min="1" max="5" step="0.1" value="1">
            <span id="rating-val" class="slider-val">All Stars</span>
        </div>
      </div>

      <div class="row" id="radius-row" style="margin-top:16px;">
        <label for="radius">Radius</label>
        <input id="radius" type="range" min="1" max="10" value="5" step="0.5">
        <span id="radiusVal">5.0 km</span>
      </div>

      <div class="legend">
        <span class="dot" style="background:#4CAF50"></span>Open / Vet<br>
        <span class="dot" style="background:#90A4AE"></span>Closed / Unknown<br>
        <span class="dot" style="background:#F44336"></span>24/7 Emergency<br>
        <span class="dot" style="background:#FFEB3B"></span>Groomer &bull; <span class="dot" style="background:#8BC34A"></span>Park (Always On) &bull; <span class="dot" style="background:#FFB300"></span>Mall &bull; <span class="dot" style="background:#fff; border: 2px solid #000;"></span>Dining
      </div>

      <div id="status">Ready to sniff out a place! 🐶</div>
  </div>
  <div class="disclaimer">
      <b>100% Unbiased:</b> MyBuddy takes <b>zero commission</b> from clinics, transport, or booking partners. All listings are equal. No one can pay to rank higher.
  </div>
</section>

<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<script>
L.TileLayer.prototype.options.referrerPolicy = 'strict-origin-when-cross-origin';
const DATA = /*DATA*/;
const SINGAPORE = [1.3521, 103.8198];
const SGT_OFFSET_MIN = 8 * 60;
const PRICE_EMAIL = "info@mybuddy.sg";

function singaporeNow() {
  const now = new Date();
  const sgt = new Date(now.getTime() + SGT_OFFSET_MIN * 60000);
  return { day: ["Sunday","Monday","Tuesday","Wednesday","Thursday","Friday","Saturday"][sgt.getUTCDay()], mins: sgt.getUTCHours() * 60 + sgt.getUTCMinutes() };
}

function statusFor(c) {
  if (c.category !== "Vet") return { key: "open", label: "OPEN" };
  if (c.ecat === "E1_CONFIRMED_24_7") return { key:"open", label:"OPEN NOW · 24/7 EMERGENCY" };
  const { day, mins } = singaporeNow();
  const iv = c.schedule[day];
  if (iv === "appointment") return { key:"appt", label:"BY APPOINTMENT ONLY" };
  if (iv === null || iv === undefined) return { key:"unknown", label:"HOURS UNCERTAIN" };
  if (!iv.length) return { key:"closed", label:"CLOSED NOW" };
  return iv.some(([a,b]) => a <= mins && mins < b) ? { key:"open",  label:"OPEN NOW" } : { key:"closed", label:"CLOSED NOW" };
}

const map = L.map("map", { center: SINGAPORE, zoom: 11, zoomControl: false });
L.control.zoom({ position: 'bottomleft' }).addTo(map);
L.tileLayer("https://www.onemap.gov.sg/maps/tiles/Default/{z}/{x}/{y}.png", { minZoom: 11, maxZoom: 19, bounds: [[1.1304, 103.6020], [1.4784, 104.0945]], attribution: '&copy; OneMap' }).addTo(map);

const centreCircle = L.circle(SINGAPORE, { radius: 5000, color:'#4CAF50', fillColor:'#4CAF50', fillOpacity:0.1, weight:3, dashArray: '5, 5', interactive:false }).addTo(map);
const centreHandle = L.marker(SINGAPORE, { icon: L.divIcon({className:'centre-handle-icon', html:'<div class="centre-handle">+</div>', iconSize:[46,46], iconAnchor:[23,23]}), draggable: true, zIndexOffset: 1000 }).addTo(map);

function moveCentre(ll) { centreCircle.setLatLng(ll); centreHandle.setLatLng(ll); updateVisibility(); }
centreHandle.on('drag', e => moveCentre(e.target.getLatLng()));
centreHandle.on('dragend', e => moveCentre(e.target.getLatLng()));

function pinIcon(c, st) {
  let colour = "open";
  let symbol = "";
  
  if (c.category === "Groomer") { colour = "groomer"; symbol = "🌼"; }
  else if (c.category === "Park") { colour = "park"; symbol = "🍃"; }
  else if (c.category === "Mall") { colour = "mall"; symbol = "🛍️"; }
  else if (c.category === "Restaurant") { colour = "dining"; symbol = "🍴"; }
  else if (c.ecat === "E1_CONFIRMED_24_7") { colour = "emergency"; symbol = "✚"; }
  else if (st.key === "closed") { colour = "closed"; }

  return L.divIcon({ className: "", html: `<div class="pin-hit"><div class="pin ${colour}">${symbol}</div></div>`, iconSize: [36, 36], iconAnchor: [18, 18], popupAnchor: [0, -12] });
}

function digits(s){ return String(s||"").replace(/\D/g,""); }
function primaryPhone(p) { return p ? digits(p.split('/')[0]) : ""; }
function esc(s){ return String(s??"").replace(/[&<>"']/g, c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c])); }

window.reportPrice = function(id, name) {
    const subject = encodeURIComponent(`Price Update for ${name} (${id})`);
    const body = encodeURIComponent(`Hi MyBuddy,\n\nI would like to update/review the pricing for ${name}.\n\nCorrect Price: \n\nThank you!`);
    window.location.href = `mailto:${PRICE_EMAIL}?subject=${subject}&body=${body}`;
}

function popupHTML(c, st) {
  const color = st.key==="open" ? "#388E3C" : "#90A4AE";
  
  const directionsHTML = (c.lat && c.lng)
    ? `<a href="https://www.google.com/maps/dir/?api=1&destination=${c.lat},${c.lng}" target="_blank" rel="noopener" style="display:flex;align-items:center;justify-content:center;gap:8px;margin:10px 0 8px;padding:12px;background:#FF9800;color:#fff !important;text-decoration:none;border-radius:12px;font-weight:800;font-size:14px;">🗺️ Get Directions</a>`
    : '';

  const phoneHTML = c.phone ? `<a class="btn-link" href="tel:${primaryPhone(c.phone)}">📞 Call ${esc(c.phone.split('/')[0].trim())}</a>` : '';

  if (c.category !== "Vet") {
    let badgeColor = "#8B4513"; 
    if (c.category === "Groomer") badgeColor = "#FFEB3B";
    if (c.category === "Mall") badgeColor = "#FFB300";
    if (c.category === "Restaurant") badgeColor = "#fff";
    
    let ratingHTML = "";
    if (c.category === "Restaurant" && c.rating > 0) {
        ratingHTML = `<div style="display:flex; align-items:center; gap:4px; margin:6px 0; font-size:13px; font-weight:800; color:#FF9800;">
            <span>⭐ ${parseFloat(c.rating).toFixed(1)}</span>
            <span style="color:#90A4AE; font-weight:600; font-size:12px;">(${c.review_count} reviews)</span>
        </div>`;
    }
    
    let txtColor = (c.category === "Restaurant") ? "#000" : (c.category === "Groomer" ? "#000" : "#fff");

    let googleLinkHTML = "";
    if (c.category === "Restaurant") {
        let gUrl = "";
        if (c.place_id) {
            gUrl = `https://www.google.com/maps/place/?q=place_id:${c.place_id}`;
        } else {
            const query = encodeURIComponent(`${c.name} ${c.address} Singapore`);
            gUrl = `https://www.google.com/maps/search/?api=1&query=${query}`;
        }
        googleLinkHTML = `<a class="btn-link" href="${gUrl}" target="_blank" rel="noopener" style="background:#4285F4; margin-top:8px; width:100%;">📖 Read Reviews & Menu</a>`;
    }

    return `<div class="popup">
      <h3>${esc(c.name)}</h3>
      <span class="badge" style="background:${badgeColor}; color:${txtColor}; ${c.category === 'Restaurant' ? 'border: 1px solid #000;' : ''}">${esc(c.category === "Park" ? "Park" : (c.category === "Mall" ? "Pet Mall" : (c.category === "Restaurant" ? "Dine with Pets" : c.category)))}</span>
      ${ratingHTML}
      <p style="font-size:13px; color:#555; margin-top:8px;">${esc(c.address)}<br>Singapore ${esc(c.postal)}</p>
      ${c.hours_text ? `<p style="font-size:13px; color:#555; margin-top:8px;">⏰ <b>Hours:</b> ${esc(c.hours_text)}</p>` : ''}
      ${directionsHTML}
      <div style="margin-top:10px; border-top:2px dashed #EEE; padding-top:10px;">
        ${phoneHTML}
        ${googleLinkHTML}
      </div>
    </div>`;
  }

  const sib = DATA.siblings[c.id] || [];
  const sibHTML = sib.length ? '<div style="border-top:2px dashed #EEE;margin-top:10px;padding-top:10px;font-size:12px;"><b style="color:#388E3C">Other branches</b><br>' + sib.map(s=>esc(s.name)).join("<br>") + '</div>' : "";
  const badges = [];
  if (c.ecat==="E1_CONFIRMED_24_7") badges.push('<span class="badge" style="background:#F44336">24/7 Emergency</span>');
  if (c.house_call) badges.push('<span class="badge" style="background:#00897B">House calls</span>');
  if (c.hospice) badges.push('<span class="badge" style="background:#5d4037">Hospice</span>');
  if (c.is_appt) badges.push('<span class="badge" style="background:#673AB7">📅 By Appointment</span>');

  let ratingHTML = c.rating > 0 ? `<div style="display:flex; align-items:center; gap:4px; margin:6px 0; font-size:13px; font-weight:800; color:#FF9800;"><span>⭐ ${parseFloat(c.rating).toFixed(1)}</span><span style="color:#90A4AE; font-weight:600; font-size:12px;">(${c.review_count} reviews)</span></div>` : "";

  let priceHTML = "";
  if (c.prices && c.prices.length > 0) {
      priceHTML = `<div style="font-weight:800; color:#388E3C; margin: 10px 0 4px; font-size:13px;">💰 Estimated Pricing</div><div class="price-list">`;
      c.prices.forEach(p => { priceHTML += `<div class="price-item">${esc(p)}</div>`; });
      priceHTML += `</div><div style="font-size:10px; color:#F44336; margin-top:4px; font-weight:700;">⚠️ Prices may be outdated. Please confirm.</div>
      <button onclick="reportPrice('${c.id}', '${esc(c.name).replace(/'/g, "\\'")}')" style="width:100%; margin-top:6px; background:#fff; border:1px solid #388E3C; color:#388E3C; padding:8px; border-radius:8px; font-weight:bold; font-family:inherit; cursor:pointer; font-size:12px;">✍️ Review / Update Price</button>`;
  }

  return `<div class="popup">
    <h3>${esc(c.name)}</h3>
    ${c.branch ? `<div style="color:#90A4AE;font-size:12px;font-weight:600;margin-bottom:6px;">${esc(c.branch)}</div>` : ''}
    ${ratingHTML}
    <p style="color:${color}; font-weight:800; font-size:13px;">${st.label}</p>
    <p style="font-size:13px; color:#555;">${esc(c.address)}<br>Singapore ${esc(c.postal)}</p>
    ${directionsHTML}
    ${badges.length ? `<div>${badges.join("")}</div>` : ''}
    ${priceHTML}
    <div style="margin-top:10px; border-top:2px dashed #EEE; padding-top:10px;">
      ${phoneHTML}
    </div>
    ${sibHTML}
  </div>`;
}

const layer = L.layerGroup().addTo(map);
const markers = []; const markersById = {}; const clinicsById = {};
let currentFilter = "all";
let currentSpecies = "all";
let currentCategory = "Vet";
let currentRating = 1.0;
let searchText = ""; let selectedId = null;

function updateVisibility() {
  const centre = centreCircle.getLatLng(); const radius = centreCircle.getRadius();
  const q = searchText.trim().toLowerCase(); const hasSearch = q.length > 0;
  let visible = 0;

  if (selectedId) {
    markers.forEach(m => {
      const on = m.clinic.id === selectedId;
      if (on) { if (!layer.hasLayer(m.marker)) layer.addLayer(m.marker); visible++; }
      else if (layer.hasLayer(m.marker)) layer.removeLayer(m.marker);
    });
    return;
  }

  markers.forEach(m => {
    const c = m.clinic, st = m.status;
    const isEmergency = c.ecat === "E1_CONFIRMED_24_7";
    
    let show = true;
    
    // PARKS ALWAYS SHOW unless searching
    if (c.category !== "Park" && currentCategory !== "All" && c.category !== currentCategory) show = false;
    
    if (show && hasSearch) {
        let match = c.name.toLowerCase().includes(q) || c.id.toLowerCase().includes(q);
        // Postal Code Search
        if (!match && /^\d{6}$/.test(q)) {
            match = (c.postal === q);
        }
        show = match;
    }
    
    if (show && !isEmergency && currentCategory === "Vet" && currentFilter === "open" && st.key !== "open") show = false;
    
    if (show && currentSpecies !== "all" && currentCategory === "Vet" && c.category === "Vet") {
      const sp = c.species || [];
      if (sp.length > 0 && !sp.includes(currentSpecies)) show = false;
    }
    
    if (show && currentCategory === "Restaurant" && currentRating > 1.0) {
       if (c.rating < currentRating) show = false;
    }
    
    if (show && !isEmergency && !hasSearch && (currentCategory === "Vet" || currentCategory === "Groomer" || currentCategory === "Restaurant")) {
      // Radius applies to Vets, Groomers, and Restaurants, but NOT Parks or Malls
      if (c.category !== "Park" && c.category !== "Mall") {
          show = map.distance(centre, m.marker.getLatLng()) <= radius;
      }
    }

    if (show) { if (!layer.hasLayer(m.marker)) layer.addLayer(m.marker); visible++; }
    else if (layer.hasLayer(m.marker)) layer.removeLayer(m.marker);
  });
  
  let catName = currentCategory;
  if (currentCategory === "Vet") catName = "Vets";
  if (currentCategory === "Groomer") catName = "Groomers";
  if (currentCategory === "Mall") catName = "Pet Malls";
  if (currentCategory === "Restaurant") catName = "Dine with Pets";

  document.getElementById("status").textContent = hasSearch ? `Search: ${visible} matching` : `Showing ${visible} ${catName}`;
}

function updateRatingUI(val) {
    let txt = "All Stars";
    if (val > 1) txt = `${parseFloat(val).toFixed(1)}+ Stars`;
    document.getElementById("rating-val").textContent = txt;
    if (document.getElementById("mobile-rating-val")) {
        document.getElementById("mobile-rating-val").textContent = txt;
    }
}

document.getElementById("rating-slider").addEventListener("input", e => {
    currentRating = parseFloat(e.target.value);
    updateRatingUI(currentRating);
    if (currentCategory !== "Restaurant") {
        document.querySelector('.cat-btn[data-cat="Restaurant"]').click();
    } else {
        updateVisibility();
    }
});

if (document.getElementById("mobile-rating-slider")) {
    document.getElementById("mobile-rating-slider").addEventListener("input", e => {
        currentRating = parseFloat(e.target.value);
        document.getElementById("rating-slider").value = currentRating;
        updateRatingUI(currentRating);
        if (currentCategory !== "Restaurant") {
            document.querySelector('.cat-btn[data-cat="Restaurant"]').click();
        } else {
            updateVisibility();
        }
    });
}

document.querySelectorAll('.cat-btn').forEach(b => {
  b.addEventListener('click', () => {
    currentCategory = b.dataset.cat;
    document.querySelectorAll('.cat-btn').forEach(x => x.classList.toggle('on', x.dataset.cat === currentCategory));
    
    const isVet = currentCategory === "Vet";
    const isGroomer = currentCategory === "Groomer";
    const isDining = currentCategory === "Restaurant";
    
    document.getElementById("species-filters").style.display = isVet ? "flex" : "none";
    
    const showRadius = isVet || isGroomer || isDining;
    document.getElementById("radius-row").style.display = showRadius ? "flex" : "none";
    if (document.getElementById("mobile-radius-slider")) {
        document.getElementById("mobile-radius-slider").style.display = showRadius ? "block" : "none";
    }
    
    document.getElementById("desktop-rating-container").style.display = isDining ? "flex" : "none";
    if (document.getElementById("mobile-rating-bar")) {
        document.getElementById("mobile-rating-bar").classList.toggle("active", isDining);
    }

    if (!isDining) {
        currentRating = 1.0;
        document.getElementById("rating-slider").value = 1;
        if (document.getElementById("mobile-rating-slider")) document.getElementById("mobile-rating-slider").value = 1;
        updateRatingUI(1);
    }

    document.getElementById("search").value = "";
    document.getElementById("mobile-search-input").value = "";
    searchText = ""; selectedId = null;
    
    updateVisibility();
  });
});

document.querySelectorAll('.sp-btn').forEach(b => {
  b.addEventListener('click', () => {
    const sp = b.dataset.sp;
    currentSpecies = (currentSpecies === sp) ? "all" : sp;
    document.querySelectorAll('.sp-btn').forEach(x => x.classList.toggle('on', x.dataset.sp === currentSpecies));
    selectedId = null; updateVisibility();
  });
});

document.getElementById("radius").addEventListener("input", e => {
  const km = parseFloat(e.target.value);
  centreCircle.setRadius(km * 1000);
  document.getElementById("radiusVal").textContent = km.toFixed(1) + " km";
  updateVisibility();
});

const mobileRadiusInput = document.getElementById("mobile-radius-input");
if (mobileRadiusInput) {
    mobileRadiusInput.addEventListener("input", e => {
        const km = parseFloat(e.target.value);
        centreCircle.setRadius(km * 1000);
        document.getElementById("radius").value = km; 
        document.getElementById("radiusVal").textContent = km.toFixed(1) + " km";
        updateVisibility();
    });
}

const searchInput = document.getElementById("search");
const directory = document.getElementById("directory");

function renderDirectory(filter) {
  const q = (filter||"").trim().toLowerCase();
  directory.innerHTML = "";
  const letters = Object.keys(DATA.directory).sort();
  letters.forEach(L => {
    const items = DATA.directory[L].filter(c => c.name.toLowerCase().includes(q) && (currentCategory === "All" || c.category === currentCategory));
    if (!items.length) return;
    const section = document.createElement("div"); section.className = "dir-letter" + (q ? " open" : "");
    const header = document.createElement("div"); header.className = "dir-header"; header.textContent = L + " (" + items.length + ")";
    header.onclick = () => section.classList.toggle("open"); section.appendChild(header);
    const list = document.createElement("div"); list.className = "dir-items";
    items.forEach(c => {
      const item = document.createElement("div"); item.className = "dir-item"; item.textContent = c.name;
      item.onclick = () => selectClinic(c.id); list.appendChild(item);
    });
    section.appendChild(list); directory.appendChild(section);
  });
}
function openDirectory() { renderDirectory(searchInput.value); directory.classList.add("open"); }
function closeDirectory() { directory.classList.remove("open"); }

searchInput.addEventListener("focus", openDirectory);
searchInput.addEventListener("click", openDirectory);
searchInput.addEventListener("input", () => {
  searchText = searchInput.value; selectedId = null;
  const q = searchText.trim().toLowerCase();
  if (q) {
    const matches = Object.values(markersById).filter(m => m.clinic.name.toLowerCase().includes(q) || m.clinic.id.toLowerCase().includes(q));
    if (matches.length === 1) selectClinic(matches[0].clinic.id, { keepInput: true });
    else updateVisibility();
  } else { updateVisibility(); }
  openDirectory();
});
document.addEventListener("click", e => { if (!e.target.closest(".search-wrap")) closeDirectory(); });

function buildMobileDropdown(filter = "") {
    const mDropdown = document.getElementById('mobile-dir-dropdown');
    const q = filter.toLowerCase();
    mDropdown.innerHTML = "";

    // DISABLE DROPDOWN FOR MALL
    if (currentCategory === "Mall") return;
    
    const letters = Object.keys(DATA.directory).sort();
    letters.forEach(letter => {
        const clinics = DATA.directory[letter].filter(c => 
            c.name.toLowerCase().includes(q) && (currentCategory === "All" || c.category === currentCategory)
        );
        if (clinics.length === 0) return;

        const header = document.createElement('div');
        header.className = 'md-letter';
        header.textContent = `${letter} (${clinics.length})`;
        
        const list = document.createElement('div');
        list.className = 'md-items';

        header.onclick = () => {
            header.classList.toggle('open');
        };

        clinics.forEach(c => {
            const item = document.createElement('div');
            item.className = 'md-item';
            item.textContent = c.name;
            item.onclick = () => {
                document.getElementById('mobile-search-input').value = c.name;
                mDropdown.classList.remove('open');
                selectClinic(c.id);
            };
            list.appendChild(item);
        });

        mDropdown.appendChild(header);
        mDropdown.appendChild(list);
    });
}

const mobileSearchInput = document.getElementById("mobile-search-input");
const mobileDropdown = document.getElementById("mobile-dir-dropdown");

if (mobileSearchInput && mobileDropdown) {
    mobileSearchInput.addEventListener("focus", () => {
        buildMobileDropdown(mobileSearchInput.value);
        mobileDropdown.classList.add("open");
    });
    mobileSearchInput.addEventListener("input", () => {
        buildMobileDropdown(mobileSearchInput.value);
    });
    document.addEventListener("click", e => {
        if (!e.target.closest(".m-search-wrap")) {
            mobileDropdown.classList.remove("open");
        }
    });
}

function clearSearch() { searchInput.value = ""; searchText = ""; selectedId = null; updateVisibility(); }
function selectClinic(id, opts) {
  const m = markersById[id]; if (!m) return;
  selectedId = id;
  if (!(opts && opts.keepInput)) { searchInput.value = m.clinic.name; searchText = m.clinic.name.toLowerCase(); }
  closeDirectory();
  updateVisibility();
  map.flyTo(m.marker.getLatLng(), 16, { duration: 0.8 }); m.marker.openPopup();
}

let unresolved = 0;
DATA.clinics.forEach(c => {
  clinicsById[c.id] = c;
  if (c.lat == null || c.lng == null) { unresolved++; return; }
  const st = statusFor(c);
  const marker = L.marker([c.lat, c.lng], { icon: pinIcon(c, st) });
  marker.bindPopup(() => popupHTML(c, statusFor(c)));
  marker.clinic = c; marker.status = st;
  markers.push({ marker, clinic: c, status: st }); markersById[c.id] = { marker, clinic: c, status: st };
  layer.addLayer(marker);
});

document.getElementById("cat-filters").dispatchEvent(new Event("click"));
document.querySelector('.cat-btn[data-cat="Vet"]').click();
if (unresolved) console.log("Unresolved coords: " + unresolved);
</script>
</body>
</html>
'''

if __name__ == "__main__":
    build()
